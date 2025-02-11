from pptx import Presentation
from pptx.util import Inches, Pt
import matplotlib.pyplot as plt
from capabilities import plot_capabilities
from challenges import plot_challenges
from actions import print_actions_pretty


# Function to save your figures as images
def save_fig_as_image(fig, filename):
    # Check if the figure is a Plotly figure by looking for the 'write_image' attribute
    if hasattr(fig, "write_image"):
        fig.write_image(filename)
    # Check if the figure is a Matplotlib figure by looking for the 'savefig' method
    elif hasattr(fig, "savefig"):
        fig.savefig(filename, bbox_inches="tight")
        plt.close(fig)
    else:
        print("The provided figure object is not recognized.")


def create_powerpoint(json_trends, json_capabilities, json_challenges, actions_json):
    # Initialize presentation
    prs = Presentation()
    # Slide titles
    slide_titles = ["Trend Radar", "Capabilities", "Challenges", "Actions"]
    image_paths = {
        "Trend Radar": "trend radar.png",
        "Capabilities": "capabilities.png",
        "Challenges": "challenges.png",
    }

    # Assuming you have functions that generate matplotlib figures for each analysis
    # You need to save those figures as images before adding them to slides
    save_fig_as_image(generate_trend_radar(json_trends), "trend radar.png")
    save_fig_as_image(plot_capabilities(json_capabilities), "capabilities.png")
    save_fig_as_image(plot_challenges(json_challenges), "challenges.png")
    # For actions, since it's text, we will add it directly

    # Add slides and populate them
    for title in slide_titles:
        slide_layout = prs.slide_layouts[2]  # Using template 2
        slide = prs.slides.add_slide(slide_layout)
        title_placeholder = slide.shapes.title
        title_placeholder.text = title

        if title != "Actions":
            img_path = image_paths[title]
            left = Inches(0.5)
            top = Inches(1.5)
            height = Inches(5.5)
            slide.shapes.add_picture(img_path, left, top, height=height)
        else:
            # Use the first BODY placeholder (index 1) for actions text
            content_placeholder = slide.placeholders[1]
            tf = content_placeholder.text_frame
            tf.text = "What we will do:"

            for (
                action
            ) in actions_json:  # Assuming actions_json is a list of action dictionaries
                p = tf.add_paragraph()
                p.text = f"{action['action']} - {action['how']} - {action['who']}"
                p.level = 0
                p.font.size = Pt(14)

    prs.save("strategy_presentation.pptx")
    print("Strategy Presentation Saved")
