# utils/md_to_docx.py
import os, subprocess, tempfile, re
from datetime import datetime
from docx import Document

TEMPLATE_PATH = "./Strategic Reports/report_template.docx"   # reference docx
OUTPUT_DIR    = "./Strategic Reports/"
REPORT_SUFFIX = "_strategic_report"

def write_markdown(md_text: str, company: str) -> str:
    out_dir = "./Strategic Reports"        # <— should match the PNG location
    os.makedirs(out_dir, exist_ok=True)
    md_path = os.path.join(out_dir, f"{company}.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md_text)
    return md_path

def fill_template_placeholders(docx_path: str, company: str):
    doc = Document(docx_path)
    today = datetime.now().strftime("%d/%m/%Y")
    for p in doc.paragraphs:
        p.text = p.text.replace("{{company_name}}", company)
        p.text = p.text.replace("{{date}}", today)
    doc.save(docx_path)

def md_to_docx(md_path: str, company: str) -> str:
    out_docx = os.path.join(OUTPUT_DIR, f"{company.capitalize()}{REPORT_SUFFIX}.docx")
    cmd = [
        "pandoc",
        md_path,
        "--from", "gfm",              # GitHub-flavoured markdown
        "--reference-doc", TEMPLATE_PATH,
        "--output", out_docx,
    ]
    subprocess.run(cmd, check=True)
    fill_template_placeholders(out_docx, company)
    return out_docx
