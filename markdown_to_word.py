# utils/markdown_to_word.py
import os
from datetime import datetime
from docx import Document
import re

TEMPLATE_PATH   = "./Strategic Reports/"
REPORTS_PATH    = "./Strategic Reports/"
TEMPLATE_NAME   = "report_template.docx"
SUFFIX          = "_strategic_report"

def md_to_docx(company: str, markdown: str) -> str:
    """
    Convert markdown text to a Word docx using the template.
    Returns the full path of the saved .docx file.
    """
    # 1. load template
    doc = Document(os.path.join(TEMPLATE_PATH, TEMPLATE_NAME))

    # 2. parse markdown
    for line in markdown.splitlines():
        line = line.rstrip()
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.strip():                       # non-empty body text
            doc.add_paragraph(line)

    # 3. replace placeholders
    today = datetime.now().strftime("%d/%m/%Y")
    for p in doc.paragraphs:
        p.text = p.text.replace("{{company_name}}", company)
        p.text = p.text.replace("{{date}}", today)

    # 4. save
    out_path = os.path.join(
        REPORTS_PATH, f"{company.capitalize()}{SUFFIX}.docx"
    )
    doc.save(out_path)
    return out_path
